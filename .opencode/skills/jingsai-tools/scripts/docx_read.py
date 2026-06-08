import sys
import json
import os

def read_docx(path):
    from docx import Document
    doc = Document(path)
    result = []
    for i, para in enumerate(doc.paragraphs):
        style = para.style.name if para.style else "Normal"
        text = para.text.strip()
        if text:
            result.append({"paragraph": i + 1, "style": style, "text": text})
    for table_idx, table in enumerate(doc.tables):
        result.append({"table": table_idx + 1, "type": "table_start"})
        for row_idx, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]
            result.append({"table": table_idx + 1, "row": row_idx + 1, "cells": cells})
        result.append({"table": table_idx + 1, "type": "table_end"})
    return result

def main():
    if len(sys.argv) < 2:
        print("Usage: python docx_read.py <path>", file=sys.stderr)
        sys.exit(1)
    path = sys.argv[1]
    if not os.path.isfile(path):
        print(f"File not found: {path}", file=sys.stderr)
        sys.exit(1)
    try:
        result = read_docx(path)
        print(json.dumps(result, ensure_ascii=False, indent=2))
    except Exception as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    main()
