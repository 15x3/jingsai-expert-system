import sys
import json
import os
import shutil
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

STYLE_MAP = {
    "标题1": "Heading 1",
    "标题2": "Heading 2",
    "标题3": "Heading 3",
    "正文": "Normal",
    "Heading 1": "Heading 1",
    "Heading 2": "Heading 2",
    "Heading 3": "Heading 3",
    "Normal": "Normal",
}

FONT_MAP = {
    "标题1": ("黑体", 16, True, WD_ALIGN_PARAGRAPH.CENTER),
    "标题2": ("黑体", 14, True, WD_ALIGN_PARAGRAPH.LEFT),
    "标题3": ("黑体", 12, True, WD_ALIGN_PARAGRAPH.LEFT),
    "正文": ("仿宋", 12, False, WD_ALIGN_PARAGRAPH.LEFT),
}

def backup_file(path):
    bak = path + ".bak"
    if not os.path.exists(bak):
        shutil.copy2(path, bak)

def find_paragraph_index(doc, target):
    for i, para in enumerate(doc.paragraphs):
        if target in para.text:
            return i
    return None

def replace_in_paragraph(para, old_text, new_text):
    if old_text not in para.text:
        return False
    for run in para.runs:
        if old_text in run.text:
            run.text = run.text.replace(old_text, new_text)
            return True
    full_text = para.text
    if old_text in full_text:
        for run in para.runs:
            run.text = ""
        if para.runs:
            para.runs[0].text = full_text.replace(old_text, new_text)
        return True
    return False

def replace_in_table(doc, old_text, new_text):
    count = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if replace_in_paragraph(para, old_text, new_text):
                        count += 1
    return count

def format_paragraph(para, style_name):
    cn_style = style_name
    if cn_style in FONT_MAP:
        font_name, font_size, bold, alignment = FONT_MAP[cn_style]
        for run in para.runs:
            run.font.name = font_name
            run.font.size = Pt(font_size)
            run.bold = bold
        para.alignment = alignment

def main():
    if len(sys.argv) < 2:
        print("Usage: python docx_edit.py <path> --action <replace|insert|format> ...", file=sys.stderr)
        sys.exit(1)

    path = sys.argv[1]
    args = sys.argv[2:]

    parsed = {}
    i = 0
    while i < len(args):
        if args[i].startswith("--"):
            key = args[i][2:]
            if i + 1 < len(args) and not args[i + 1].startswith("--"):
                parsed[key] = args[i + 1]
                i += 2
            else:
                parsed[key] = True
                i += 1
        else:
            i += 1

    action = parsed.get("action", "replace")
    target = parsed.get("target", "")
    replacement = parsed.get("replacement", "")
    text = parsed.get("text", "")
    style = parsed.get("style", "正文")

    if not os.path.isfile(path):
        print(f"File not found: {path}", file=sys.stderr)
        sys.exit(1)

    backup_file(path)
    doc = Document(path)
    changes = 0

    if action == "replace":
        if not target or not replacement:
            print("Error: --target and --replacement required for replace", file=sys.stderr)
            sys.exit(1)
        for para in doc.paragraphs:
            if replace_in_paragraph(para, target, replacement):
                changes += 1
        changes += replace_in_table(doc, target, replacement)

    elif action == "insert":
        if not target:
            print("Error: --target required for insert", file=sys.stderr)
            sys.exit(1)
        insert_text = text if text else replacement
        idx = find_paragraph_index(doc, target)
        if idx is not None:
            new_para = doc.paragraphs[idx]._element
            from docx.oxml.ns import qn
            from lxml import etree
            p = etree.SubElement(new_para.getparent(), qn('w:p'))
            r = etree.SubElement(p, qn('w:r'))
            t = etree.SubElement(r, qn('w:t'))
            t.text = insert_text
            new_para.addnext(p)
            changes = 1
        else:
            doc.add_paragraph(insert_text)
            changes = 1

    elif action == "format":
        if not target:
            print("Error: --target required for format", file=sys.stderr)
            sys.exit(1)
        for para in doc.paragraphs:
            if target in para.text:
                format_paragraph(para, style)
                changes += 1

    doc.save(path)
    result = {
        "status": "ok",
        "action": action,
        "changes": changes,
        "file": path,
        "backup": path + ".bak"
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))

if __name__ == "__main__":
    main()
