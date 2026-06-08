import sys
import json
import os

def read_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    result = []
    for slide_idx, slide in enumerate(prs.slides):
        slide_info = {"slide": slide_idx + 1, "shapes": []}
        for shape in slide.shapes:
            shape_info = {
                "name": shape.name,
                "type": str(shape.shape_type),
            }
            if shape.has_text_frame:
                texts = []
                for para in shape.text_frame.paragraphs:
                    para_text = para.text.strip()
                    if para_text:
                        texts.append(para_text)
                if texts:
                    shape_info["text"] = texts
            if hasattr(shape, "image"):
                shape_info["has_image"] = True
            slide_info["shapes"].append(shape_info)
        result.append(slide_info)
    return result

def main():
    if len(sys.argv) < 2:
        print("Usage: python pptx_read.py <path>", file=sys.stderr)
        sys.exit(1)
    path = sys.argv[1]
    if not os.path.isfile(path):
        print(f"File not found: {path}", file=sys.stderr)
        sys.exit(1)
    try:
        result = read_pptx(path)
        print(json.dumps(result, ensure_ascii=False, indent=2))
    except Exception as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    main()
