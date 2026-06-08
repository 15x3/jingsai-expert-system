import sys
import json
import os
import shutil

def edit_pptx(path, slide_num, action, target=None, replacement=None, text=None, position=None):
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.enum.text import PP_ALIGN

    prs = Presentation(path)
    if slide_num < 1 or slide_num > len(prs.slides):
        return {"status": "error", "message": f"Slide {slide_num} not found (total: {len(prs.slides)})"}

    slide = prs.slides[slide_num - 1]
    changes = 0

    if action == "replace":
        if not target or replacement is None:
            return {"status": "error", "message": "--target and --replacement required"}
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if target in run.text:
                            run.text = run.text.replace(target, replacement)
                            changes += 1

    elif action == "add-text":
        if not text:
            return {"status": "error", "message": "--text required for add-text"}
        pos = position or {"left": 100, "top": 100, "width": 400, "height": 50}
        left = Emu(pos.get("left", 100))
        top = Emu(pos.get("top", 100))
        width = Emu(pos.get("width", 400))
        height = Emu(pos.get("height", 50))
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.text = text
        changes = 1

    prs.save(path)
    return {"status": "ok", "action": action, "slide": slide_num, "changes": changes}

def main():
    if len(sys.argv) < 2:
        print("Usage: python pptx_edit.py <path> --slide N --action <replace|add-text> ...", file=sys.stderr)
        sys.exit(1)

    path = sys.argv[1]
    args = sys.argv[2:]

    parsed = {}
    i = 0
    while i < len(args):
        if args[i].startswith("--"):
            key = args[i][2:]
            if i + 1 < len(args) and not args[i + 1].startswith("--"):
                parsed[key] = args[i + 1]
                i += 2
            else:
                parsed[key] = True
                i += 1
        else:
            i += 1

    slide_num = int(parsed.get("slide", 1))
    action = parsed.get("action", "replace")
    target = parsed.get("target", "")
    replacement = parsed.get("replacement", "")
    text = parsed.get("text", "")
    position_str = parsed.get("position", "")

    position = None
    if position_str:
        try:
            position = json.loads(position_str)
        except:
            position = None

    if not os.path.isfile(path):
        print(f"File not found: {path}", file=sys.stderr)
        sys.exit(1)

    bak = path + ".bak"
    if not os.path.exists(bak):
        shutil.copy2(path, bak)

    result = edit_pptx(path, slide_num, action, target, replacement, text, position)
    print(json.dumps(result, ensure_ascii=False, indent=2))

if __name__ == "__main__":
    main()
